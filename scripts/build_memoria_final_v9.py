"""Integra el bloque final TSMAI V9 en la memoria académica revisada."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "docs" / "Memoria_TFM_Movilidad_Sostenible_Mallorca_final_revisada.docx"
OUTPUT = ROOT / "docs" / "Memoria_TFM_Movilidad_Sostenible_Mallorca_final_V9.docx"


def set_text(paragraph: Paragraph, text: str) -> Paragraph:
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    paragraph.add_run(text)
    return paragraph


def after(paragraph: Paragraph, text: str, style: str = "normal") -> Paragraph:
    new_xml = OxmlElement("w:p")
    paragraph._p.addnext(new_xml)
    new = Paragraph(new_xml, paragraph._parent)
    new.style = style
    new.add_run(text)
    return new


def find(doc: Document, beginning: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith(beginning):
            return paragraph
    raise ValueError(f"No se encontró el párrafo: {beginning!r}")


def main() -> None:
    shutil.copyfile(INPUT, OUTPUT)
    doc = Document(OUTPUT)
    title = find(doc, "7.10.")
    title.style = "Heading 2"
    set_text(title, "7.10 Índice TSMAI V9 y análisis estacional")
    v9_results = find(doc, "El TSMAI v2")
    set_text(v9_results, "El resultado vigente es TSMAI V9 por alojamiento. Integra siete componentes de transporte público, caminabilidad y bicicleta: proximidad y servicio de transporte público, TRANSIT temporal, evidencia y rutas BICYCLE, y evidencia y rutas WALK. El universo actual contiene 1.461 alojamientos oficiales con geometría válida. La puntuación media es 0,625 (desviación estándar 0,179): 719 alojamientos se clasifican en nivel favorable, 570 en nivel intermedio y 172 en prioridad de mejora. Los pesos se presentan como decisiones analíticas transparentes, no como coeficientes aprendidos.")
    context = find(doc, "Como contexto ambiental")
    set_text(context, "El componente TRANSIT se obtiene como media simple de dos campañas de cuatro escenarios cada una, dentro del mismo GTFS versionado. Su puntuación media es 0,4291 en verano, 0,0137 en invierno y 0,2214 en la media estacional; 793 alojamientos empeoran y 668 no cambian bajo la regla configurada. Este resultado describe disponibilidad programada en las fechas seleccionadas y no mide puntualidad, ocupación, demanda ni un promedio anual.")
    anchor = after(context, "7.11 Capas territoriales actuales", "Heading 2")
    anchor = after(anchor, "El mapa P7 vincula el mismo universo vigente con 772 paradas GTFS TIB, 4.487 destinos OSM temáticos y 1.421 evidencias ciclistas OSM. Como cribado territorial, 885 alojamientos están a 0–400 m de una parada, 267 a 400–800 m, 42 a 800–1.200 m y 267 a más de 1.200 m. Esta proximidad es euclídea y no sustituye el itinerario peatonal por red. Las cifras históricas de 1.463 alojamientos y 1.423 destinos nombrados se conservan como línea base, pero no se mezclan con el resultado V9 porque proceden de snapshots y reglas de selección distintos.")
    anchor = after(anchor, "7.12 Sensibilidad de las ponderaciones", "Heading 2")
    after(anchor, "La sensibilidad recalcula TSMAI V9 con cinco escenarios deterministas de pesos y los mismos siete componentes. En 625 alojamientos el nivel se mantiene en todos los escenarios y en 836 cambia al menos una vez. De los 147 alojamientos que forman aproximadamente un decil, 110 aparecen en el decil superior en al menos cuatro escenarios. Las correlaciones de Spearman entre rankings de alojamiento oscilan entre 0,855 y 0,995; en municipios, entre 0,916 y 0,996. La estabilidad relativa permite identificar candidaturas robustas para revisión, pero no valida causalmente el índice ni prescribe inversiones.")
    limitations = find(doc, "9. Limitaciones que")
    limitations.style = "Heading 2"
    set_text(limitations, "8.1 Limitaciones y condiciones de interpretación")
    conclusion = find(doc, "El TFM demuestra")
    set_text(conclusion, "El TFM demuestra la viabilidad de construir con herramientas abiertas un sistema reproducible de accesibilidad turística y movilidad sostenible en Mallorca. Su resultado final, TSMAI V9, combina evidencia espacial, horarios GTFS y enrutamiento OTP en un índice individual explicable, con fuentes, pesos y escenarios trazables. El resultado se orienta a priorizar investigación y validación técnica, no a certificar accesibilidad ni a sustituir una decisión pública.")
    last = find(doc, "El resultado principal no")
    set_text(last, "El resultado principal es un método actualizable que permite revisar dónde se concentran limitaciones de primera y última milla, conectividad peatonal, movilidad ciclista y servicio programado. Los 172 alojamientos en prioridad de mejora y las prioridades robustas son una cartera para contraste de campo y consulta institucional. Antes de proponer obras o estimar cambio modal se requieren datos de demanda, seguridad por tramo, continuidad física y accesibilidad universal verificadas.")
    doc.core_properties.title = "Generador de mapas de accesibilidad y movilidad sostenible en Mallorca"
    doc.core_properties.subject = "Memoria académica integrada con resultados TSMAI V9"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
